"""
Data loader for the MDF production report (LINE 1 and LINE 2).

LINE 1 (verified on the real file "ป 69 L1"):
  - sheet names look like "ก.ย. 69"; other sheets (Item, % P-U, Batch Glue) are skipped
  - header rows 2-3, data starts at row 4, the row labelled TOTAL is skipped
  - A=Date, I=Thickness (mm), J=Board type, Y=Reject (LG m3 of finished board), Z=Production (m3 of finished board)
  - Resin: 1 resin = 5 columns (kg, adjust, NO/TOTAL, %Glue, kg/m3) -> only the 3rd column (net used) is used
      -> the first resin column starts right after the "wood (kg)" column, so its position changes by month
      -> resin names are read from header row 3 of the net-used column (they also change by month)
      -> the search stops at the column named "สีเขียวผสมน้ำ" (green colour); it and everything to its right is ignored

LINE 2 workbook structure (verified on the real file):
  - 1 sheet = 1 month (Thai month names as sheet names); other sheets are skipped
  - header rows 3-4, data starts at row 5, the summary row labelled TOTAL is skipped
  - A=Date, I=Thickness (mm), J=Board type, Y=Reject after press (m3), Z=Production (m3)
  - From column AI to the right: 1 resin = 3 columns (kgs, adjust, TT)
      -> only the 3rd column (TT = net used) is used
      -> non-resin columns (fiber / CHIP / wood) are skipped
      -> the search stops at the column named "สีแดง" (red); red and everything to its right is ignored
"""
import csv
import datetime as dt
import io
import re
from collections import Counter

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CI

THAI_MONTHS = {  # sheet names in the source file
    "ม.ค.": 1, "ก.พ.": 2, "มี.ค.": 3, "เม.ย.": 4, "พ.ค.": 5, "มิ.ย.": 6,
    "ก.ค.": 7, "ส.ค.": 8, "ก.ย.": 9, "ต.ค.": 10, "พ.ย.": 11, "ธ.ค.": 12,
}
MONTH_LABEL = {1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "Jun",
               7: "Jul", 8: "Aug", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dec"}

# Column layout per production line
LINE_CONFIG = {
    "LINE 2": {
        "hdr_rows": (3, 4),
        "first_data_row": 5,
        "col_date": "A",
        "col_thick": "I",
        "col_board": "J",
        "col_reject": "Y",
        "col_prod": "Z",
        "col_glue_start": "AI",
        "glue_cols_per_resin": 3,          # kgs, adjust, TT
        "glue_name_from": "header",        # name = header text of the first column of the resin
        "stop_keyword": "สีแดง",
        "skip_prefix": ("fiber", "chip", "wood"),
    },
    "LINE 1": {
        "hdr_rows": (2, 3),
        "first_data_row": 4,
        "col_date": "A",
        "col_thick": "I",
        "col_board": "J",
        "col_reject": "Y",
        "col_prod": "Z",
        "col_glue_start": "AI",            # fallback only; the real start is found from the "wood (kg)" column
        "glue_start_after": ("wood", "(kg)"),  # (header row 1 text, header row 2 text) of the last non-resin column
        "glue_cols_per_resin": 5,          # kg, adjust, NO/TOTAL, %Glue, kg/m3
        "glue_name_from": "net_col_row2",  # name = header row 2 (hdr_rows[1]) of the net-used column
        "stop_keyword": "สีเขียวผสมน้ำ",
        "skip_prefix": (),
    },
}

# Thickness ranges for the pie chart: (label, low, high, low_is_inclusive)
THICK_BANDS = [
    ("2.2 – 4.0 mm", 2.2, 4.0, True),
    ("> 4.0 – 9.0 mm", 4.0, 9.0, False),
    ("> 9.0 – 15.0 mm", 9.0, 15.0, False),
    ("> 15.0 – 18.0 mm", 15.0, 18.0, False),
    ("> 18.0 – 25.0 mm", 18.0, 25.0, False),
]

GLUE_PREFIX = "glue::"
UNKNOWN_BOARD = "(Unspecified)"


def _cell(row, idx):
    return row[idx] if idx < len(row) else None


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _header(rows, cfg, idx):
    parts = []
    for r in cfg["hdr_rows"]:
        v = _cell(rows[r - 1], idx)
        if v not in (None, ""):
            parts.append(str(v).strip())
    return " ".join(parts)


def _resin_name(header_text):
    return re.sub(r"\s*\(kgs\)\s*", " ", header_text).strip()


def _month_from_sheet(name):
    """'ม.ค. ' (LINE 2) and 'ก.ย. 69' (LINE 1) -> month number; other sheet names -> None."""
    name = name.strip()
    if name in THAI_MONTHS:
        return THAI_MONTHS[name]
    parts = name.split()
    return THAI_MONTHS.get(parts[0]) if parts else None


def _norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v not in (None, "") else ""


def _find_glue_groups(rows, cfg, i_start, width):
    """Return [(resin name, idx first col, idx adjust col, idx net-used col), ...] for one sheet."""
    per = cfg["glue_cols_per_resin"]
    marker = cfg.get("glue_start_after")
    if marker:  # LINE 1: the resin block starts right after the last non-resin column (varies by month)
        r1, r2 = cfg["hdr_rows"]
        for c in range(max(i_start - 3, 0), width):
            if _norm(_cell(rows[r1 - 1], c)) == marker[0] and _norm(_cell(rows[r2 - 1], c)) == marker[1]:
                i_start = c + 1
                break
    groups = []
    c = i_start
    while c < width:
        h = _header(rows, cfg, c)
        if cfg["stop_keyword"] in h:
            break
        if not h or h.lower().startswith(cfg["skip_prefix"]):
            c += 1
            continue
        if cfg["glue_name_from"] == "net_col_row2":
            nm = _cell(rows[cfg["hdr_rows"][1] - 1], c + 2)
            name = str(nm).strip() if nm not in (None, "") else _resin_name(h)
        else:
            name = _resin_name(h)
        groups.append((name, c, c + 1, c + 2))
        c += per
    return groups


def parse_workbook(file_bytes, line="LINE 2"):
    """Return (df, resin_names, notes). One df row = one data row of the workbook."""
    cfg = LINE_CONFIG[line]
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    i_date, i_board = CI(cfg["col_date"]) - 1, CI(cfg["col_board"]) - 1
    i_rej, i_prod = CI(cfg["col_reject"]) - 1, CI(cfg["col_prod"]) - 1
    i_thk = CI(cfg["col_thick"]) - 1
    i_start = CI(cfg["col_glue_start"]) - 1

    frames, notes, resin_order = [], [], []

    for ws in wb.worksheets:
        month = _month_from_sheet(ws.title)
        if month is None:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < cfg["first_data_row"]:
            continue

        width = max(len(r) for r in rows)
        groups = _find_glue_groups(rows, cfg, i_start, width)  # (resin name, idx kgs, idx adjust, idx TT)

        recs = []
        for r in rows[cfg["first_data_row"] - 1:]:
            d = _cell(r, i_date)
            if not isinstance(d, (dt.datetime, dt.date)):
                continue  # skip TOTAL / blank rows
            thk = _cell(r, i_thk)
            rec = {
                "date": pd.Timestamp(d).normalize(),
                "thick": float(thk) if isinstance(thk, (int, float)) and not isinstance(thk, bool) else float("nan"),
                "board": (str(_cell(r, i_board)).strip()
                          if _cell(r, i_board) not in (None, "") else UNKNOWN_BOARD),
                "reject_m3": _num(_cell(r, i_rej)),
                "prod_m3": _num(_cell(r, i_prod)),
            }
            for name, _, _, i_tt in groups:
                rec[GLUE_PREFIX + name] = rec.get(GLUE_PREFIX + name, 0.0) + _num(_cell(r, i_tt))
            recs.append(rec)
        if not recs:
            continue
        df = pd.DataFrame(recs)

        # drop empty shift rows (no board type and no quantities)
        gcols = [GLUE_PREFIX + n for n, *_ in groups]
        empty = (df["board"] == UNKNOWN_BOARD) & (df[["prod_m3", "reject_m3"] + gcols].sum(axis=1) == 0)
        df = df[~empty].reset_index(drop=True)
        if df.empty:
            continue

        # fix mistyped years (e.g. 2029) using the most common year of the sheet
        year = Counter(df["date"].dt.year).most_common(1)[0][0]
        bad = df["date"].dt.year != year
        if bad.any():
            fixed = []
            for d in df.loc[bad, "date"]:
                try:
                    fixed.append(d.replace(year=year))
                except ValueError:
                    fixed.append(d.replace(year=year, day=28))
            df.loc[bad, "date"] = fixed
            notes.append(f"Sheet {MONTH_LABEL[month]}: corrected {int(bad.sum())} row(s) with an abnormal year -> {year}")
        df["month"] = month

        # sanity check: TT should equal kgs + adjust
        data_rows = [r for r in rows[cfg["first_data_row"] - 1:]
                     if isinstance(_cell(r, i_date), (dt.datetime, dt.date))]
        for name, i1, i2, i3 in groups:
            s1 = sum(_num(_cell(r, i1)) for r in data_rows)
            s2 = sum(_num(_cell(r, i2)) for r in data_rows)
            s3 = sum(_num(_cell(r, i3)) for r in data_rows)
            if abs(s3 - (s1 + s2)) > max(1.0, 0.005 * abs(s3)):
                notes.append(f"Sheet {MONTH_LABEL[month]}: resin {name} TT differs from kgs + adjust "
                             f"(TT={s3:,.0f}, kgs+adjust={s1 + s2:,.0f}); the TT value is used")
        for name, *_ in groups:
            if name not in resin_order:
                resin_order.append(name)
        frames.append(df)

    if not frames:
        raise ValueError("No monthly sheets (Jan-Dec) with data were found in this file.")

    out = pd.concat(frames, ignore_index=True)
    glue_cols = [GLUE_PREFIX + n for n in resin_order]
    out[glue_cols] = out[glue_cols].fillna(0.0)
    out["resin_kg"] = out[glue_cols].sum(axis=1)
    out = out.sort_values("date", kind="stable").reset_index(drop=True)
    return out, resin_order, notes


# ---------- summaries ----------
def filter_rows(df, months=None, date_range=None):
    m = pd.Series(True, index=df.index)
    if months is not None:
        m &= df["month"].isin(months)
    if date_range is not None:
        m &= (df["date"] >= pd.Timestamp(date_range[0])) & (df["date"] <= pd.Timestamp(date_range[1]))
    return df[m]


def kpis(d):
    prod, rej = d["prod_m3"].sum(), d["reject_m3"].sum()
    return {"prod": float(prod), "reject": float(rej),
            "reject_pct": float(rej / prod * 100) if prod else None}


def glue_table(d, resin_names):
    """Net resin usage (TT) per resin code, largest first (zero-usage resins are dropped)."""
    rows = [(n, d[GLUE_PREFIX + n].sum()) for n in resin_names]
    t = pd.DataFrame(rows, columns=["Resin", "Net used (kg)"])
    total = float(t["Net used (kg)"].sum())
    t = t[t["Net used (kg)"] > 0].sort_values("Net used (kg)", ascending=False)
    t["Share (%)"] = t["Net used (kg)"] / total * 100 if total else 0
    return t.reset_index(drop=True), total


# ---------- board type grouping ----------
def read_group_map_csv(path):
    """Read the Board type -> Group table from a CSV (columns: Board type, Group)."""
    m = {}
    try:
        with open(path, encoding="utf-8-sig", newline="") as f:
            for row in csv.reader(f):
                if len(row) >= 2 and row[0].strip() and row[0].strip() != "Board type":
                    if row[1].strip():
                        m[row[0].strip()] = row[1].strip()
    except FileNotFoundError:
        pass
    return m


def read_group_map_xlsx(file_bytes):
    """Read the grouping table from an Excel file (column A = Board type, B = Group)."""
    ws = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True).worksheets[0]
    m = {}
    for a, b, *_ in ws.iter_rows(values_only=True):
        a = str(a).strip() if a is not None else ""
        b = str(b).strip() if b is not None else ""
        if not a or a.startswith(("Board type", "รวมทั้งหมด")):
            continue
        if b and not b.startswith("กลุ่ม"):
            m[a] = b
    return m


def apply_groups(df, mapping):
    """Add a 'group' column; unmapped names keep their own name. Returns (df, unmapped names)."""
    df = df.copy()
    df["group"] = df["board"].map(mapping)
    missing = sorted(df.loc[df["group"].isna() & (df["board"] != UNKNOWN_BOARD), "board"].unique())
    df["group"] = df["group"].fillna(df["board"])
    return df, missing


def board_table(d, by="board"):
    g = d.groupby(by, as_index=False).agg(m3=("prod_m3", "sum"), resin=("resin_kg", "sum"))
    g = g.sort_values("m3", ascending=False).reset_index(drop=True)
    tot = g["m3"].sum()
    g["share"] = g["m3"] / tot * 100 if tot else 0
    g["cons"] = g.apply(lambda r: r["resin"] / r["m3"] if r["m3"] else None, axis=1)
    g.insert(0, "Rank", range(1, len(g) + 1))
    g.columns = ["Rank", "Board Type Group" if by == "group" else "Board Type",
                 "Production (m³)", "Resin (kg)", "Share (%)", "Consumption (kg/m³)"]
    return g


def top_n_with_others(bt, n=8):
    """Labels/values for a pie chart: the top n rows plus one 'Others' slice."""
    label_col = bt.columns[1]
    top = bt.head(n)
    labels = list(top[label_col])
    values = [float(v) for v in top["Production (m³)"]]
    rest = float(bt["Production (m³)"].iloc[n:].sum())
    if rest > 0:
        labels.append(f"Others ({len(bt) - n})")
        values.append(rest)
    return labels, values


def thickness_table(d):
    """Production per thickness range. Returns (table, info) where info describes volume outside all ranges."""
    t = d["thick"].round(3)
    assigned = pd.Series(False, index=d.index)
    rows = []
    for label, lo, hi, lo_inc in THICK_BANDS:
        m = ((t >= lo) if lo_inc else (t > lo)) & (t <= hi)
        rows.append((label, float(d.loc[m, "prod_m3"].sum())))
        assigned |= m
    tbl = pd.DataFrame(rows, columns=["Thickness Range", "Production (m³)"])
    classified = float(tbl["Production (m³)"].sum())
    tbl["Share (%)"] = tbl["Production (m³)"] / classified * 100 if classified else 0.0
    out = d.loc[~assigned & (d["prod_m3"] > 0)]
    info = {"classified": classified,
            "outside_m3": float(out["prod_m3"].sum()),
            "outside_values": sorted({round(float(v), 2) for v in out["thick"].dropna()})}
    return tbl, info


def product_table(d, by="group"):
    """Production and resin consumption per Product = thickness + board type (group)."""
    x = d[(d["prod_m3"] > 0) & d["thick"].notna()].copy()
    x["_t"] = x["thick"].round(3)
    g = x.groupby(["_t", by], as_index=False).agg(m3=("prod_m3", "sum"), resin=("resin_kg", "sum"))
    g["Product"] = g.apply(lambda r: f"{r['_t']:g} {r[by]}", axis=1)
    g = g.sort_values(["m3", "_t"], ascending=[False, True]).reset_index(drop=True)
    g["cons"] = g.apply(lambda r: r["resin"] / r["m3"] if r["m3"] else None, axis=1)
    g.insert(0, "Rank", range(1, len(g) + 1))
    g = g[["Rank", "Product", "m3", "cons"]]
    g.columns = ["Rank", "Product", "Production (m³)", "Consumption (kg/m³)"]
    return g


def product_monthly(df, label, by="group"):
    """Monthly production and resin consumption of ONE product (label like '18 UF') over every month in df.

    Months in which the product was not produced are kept (NaN consumption) so the chart shows a gap.
    """
    x = df[(df["prod_m3"] > 0) & df["thick"].notna()].copy()
    x["_p"] = x["thick"].round(3).map(lambda t: f"{t:g}") + " " + x[by].astype(str)
    x = x[x["_p"] == label]
    months = sorted(int(m) for m in df["month"].unique())
    g = x.groupby("month").agg(m3=("prod_m3", "sum"), resin=("resin_kg", "sum")).reindex(months)
    g.index.name = "month"
    g["cons"] = g["resin"] / g["m3"].where(g["m3"] > 0)
    g = g.reset_index()
    g.insert(1, "Month", g["month"].map(MONTH_LABEL))
    g = g.drop(columns="month")
    g.columns = ["Month", "Production (m³)", "Resin (kg)", "Consumption (kg/m³)"]
    return g
